"""
Tests for Celery Worker Tasks (Milestone 3).
"""

import hashlib
import io
import tempfile
import uuid
from pathlib import Path

import openpyxl
import pytest
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.db.models import Base
from app.db.repositories import DatasetRepository, UploadRepository
from app.storage import LocalStorageBackend
from app.worker.tasks import normalize_workbook_task, scan_workbook_task


def _make_xlsx(sheets: dict) -> bytes:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for sheet_name, rows in sheets.items():
        ws = wb.create_sheet(title=sheet_name)
        for r_idx, row in enumerate(rows, start=1):
            for c_idx, val in enumerate(row, start=1):
                ws.cell(row=r_idx, column=c_idx, value=val)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.fixture
def worker_context(monkeypatch):
    """Sets up in-memory DB and temp storage for worker tasks."""
    engine = create_engine("sqlite:///:memory:", echo=False)
    Base.metadata.create_all(bind=engine)
    TestingSession = sessionmaker(bind=engine)

    from contextlib import contextmanager

    @contextmanager
    def mock_db_context():
        session = TestingSession()
        try:
            yield session
            session.commit()
        except Exception:
            session.rollback()
            raise
        finally:
            session.close()

    temp_dir = Path(tempfile.mkdtemp())
    test_storage = LocalStorageBackend(root_dir=temp_dir)

    import app.worker.tasks as tasks_mod

    monkeypatch.setattr(tasks_mod, "get_db_context", mock_db_context)
    monkeypatch.setattr(tasks_mod, "storage", test_storage)
    monkeypatch.setattr(tasks_mod, "get_storage_backend", lambda: test_storage)

    yield {
        "engine": engine,
        "sessionmaker": TestingSession,
        "storage": test_storage,
    }


class TestScanWorkbookTask:

    def test_scan_task_success_and_dataset_creation(self, worker_context):
        storage = worker_context["storage"]
        Session = worker_context["sessionmaker"]

        # 1. Create file and upload record
        xlsx_bytes = _make_xlsx(
            {
                "Defects": [
                    ["Defect", "Jan-26", "Feb-26"],
                    ["Scratch", 10, 15],
                    ["Dent", 5, 12],
                ]
            }
        )
        upload_id = uuid.uuid4()
        storage_key = f"uploads/{upload_id}/original.xlsx"
        storage.put_object(storage_key, xlsx_bytes)

        db = Session()
        upload_repo = UploadRepository(db)
        upload = upload_repo.create(
            upload_id=upload_id,
            original_filename="Defects_2026.xlsx",
            sanitized_filename="Defects_2026.xlsx",
            storage_key=storage_key,
            file_size_bytes=len(xlsx_bytes),
            file_type="xlsx",
            sha256_hash=hashlib.sha256(xlsx_bytes).hexdigest(),
        )
        db.close()

        # 2. Execute task
        result = scan_workbook_task(str(upload_id), storage_key)
        assert result["status"] == "success"
        assert result["dataset_count"] >= 1

        # 3. Verify upload status in DB
        db = Session()
        upload_repo = UploadRepository(db)
        dataset_repo = DatasetRepository(db)

        updated_upload = upload_repo.get(upload_id)
        assert updated_upload.status in ("needs_mapping", "ready")
        assert updated_upload.progress == 100

        datasets = dataset_repo.list_by_upload(upload_id)
        assert len(datasets) >= 1
        assert datasets[0].sheet_name == "Defects"
        db.close()

    def test_scan_task_retry_idempotence(self, worker_context):
        """Re-running the scan task must not multiply dataset rows."""
        storage = worker_context["storage"]
        Session = worker_context["sessionmaker"]

        xlsx_bytes = _make_xlsx({"S1": [["A", "B"], [1, 2]]})
        upload_id = uuid.uuid4()
        storage_key = f"uploads/{upload_id}/original.xlsx"
        storage.put_object(storage_key, xlsx_bytes)

        db = Session()
        UploadRepository(db).create(
            upload_id=upload_id,
            original_filename="S1.xlsx",
            sanitized_filename="S1.xlsx",
            storage_key=storage_key,
            file_size_bytes=len(xlsx_bytes),
            file_type="xlsx",
            sha256_hash=hashlib.sha256(xlsx_bytes).hexdigest(),
        )
        db.close()

        # Run task twice
        scan_workbook_task(str(upload_id), storage_key)
        scan_workbook_task(str(upload_id), storage_key)

        db = Session()
        datasets = DatasetRepository(db).list_by_upload(upload_id)
        # Exactly 1 dataset created, not 2
        assert len(datasets) == 1
        db.close()

    def test_scan_task_failure_records_error(self, worker_context):
        storage = worker_context["storage"]
        Session = worker_context["sessionmaker"]

        upload_id = uuid.uuid4()
        storage_key = f"uploads/{upload_id}/corrupted.xlsx"
        storage.put_object(storage_key, b"NOT_A_VALID_XLSX")

        db = Session()
        UploadRepository(db).create(
            upload_id=upload_id,
            original_filename="bad.xlsx",
            sanitized_filename="bad.xlsx",
            storage_key=storage_key,
            file_size_bytes=16,
            file_type="xlsx",
            sha256_hash="0000",
        )
        db.close()

        result = scan_workbook_task(str(upload_id), storage_key)
        assert result["status"] == "failed"

        db = Session()
        upload = UploadRepository(db).get(upload_id)
        assert upload.status == "failed"
        assert upload.error_code == "SCAN_ERROR"
        assert upload.error_message is not None
        db.close()


class TestNormalizeWorkbookTask:

    def test_normalize_task_creates_parquet(self, worker_context):
        storage = worker_context["storage"]
        Session = worker_context["sessionmaker"]

        xlsx_bytes = _make_xlsx(
            {
                "Production": [
                    ["Model", "Jan-26", "Feb-26"],
                    ["Sedan", 100, 150],
                    ["SUV", 200, 250],
                ]
            }
        )
        upload_id = uuid.uuid4()
        storage_key = f"uploads/{upload_id}/original.xlsx"
        storage.put_object(storage_key, xlsx_bytes)

        db = Session()
        upload = UploadRepository(db).create(
            upload_id=upload_id,
            original_filename="Prod.xlsx",
            sanitized_filename="Prod.xlsx",
            storage_key=storage_key,
            file_size_bytes=len(xlsx_bytes),
            file_type="xlsx",
            sha256_hash="hash",
        )
        dataset = DatasetRepository(db).create(
            upload_id=upload_id,
            sheet_name="Production",
            cell_range="A1:C3",
            kind="wide",
            column_metadata=[],
        )
        db.close()

        template_payload = {
            "id": "tpl-norm",
            "name": "Norm Template",
            "version": 1,
            "created_at": "2026-10-25T00:00:00Z",
            "updated_at": "2026-10-25T00:00:00Z",
            "blocks": [
                {
                    "block_index": 0,
                    "sheet": "Production",
                    "range": "A1:C3",
                    "included": True,
                    "unpivot": {
                        "enabled": True,
                        "id_columns": ["Model"],
                        "value_columns": ["Jan-26", "Feb-26"],
                        "variable_name": "Month",
                        "value_name": "Units",
                        "measure_name": "Production",
                    },
                }
            ],
            "global_warnings_acknowledged": True,
        }

        result = normalize_workbook_task(str(upload_id), storage_key, template_payload)
        assert result["status"] == "success"
        assert len(result["parquet_datasets"]) >= 1

        # Verify Parquet file exists in storage
        parquet_key = result["parquet_datasets"][0]
        assert storage.object_exists(parquet_key) is True

        # Verify dataset in DB has parquet_storage_key
        db = Session()
        updated_ds = DatasetRepository(db).get(dataset.id)
        assert updated_ds.parquet_storage_key is not None
        assert updated_ds.normalization_status == "normalized"
        db.close()
