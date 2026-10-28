"""
Unit and Integration Tests for Milestone 2: Mapping and Normalisation Engine.
"""

from __future__ import annotations

import io
import json
import os
import shutil
import tempfile
from pathlib import Path
from typing import Any, Dict, List

import openpyxl
import pandas as pd
import pytest
from fastapi.testclient import TestClient

from app.auth import AuthenticatedUser, get_current_user
from app.ingestion.mapping import (
    AmbiguousRowDecision,
    BlockMapping,
    CellColourStatusMap,
    ColumnRoleMapping,
    MappingTemplate,
    MappingTemplateSummary,
    NormalizationResult,
    SourceOfTruthConfirmation,
    TemplateStorage,
    UnpivotConfig,
    ValueMergeDecision,
    normalize_workbook,
)
from app.main import app

client = TestClient(app)


# ---------------------------------------------------------------------------
# Test Helpers
# ---------------------------------------------------------------------------

def _make_xlsx(sheets: Dict[str, List[List[Any]]], cell_colors: Dict[str, Dict[str, str]] = None) -> bytes:
    """Create an in-memory XLSX workbook with optional fill colors."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # Remove default sheet

    for sheet_name, rows in sheets.items():
        ws = wb.create_sheet(title=sheet_name)
        for r_idx, row in enumerate(rows, start=1):
            for c_idx, val in enumerate(row, start=1):
                cell = ws.cell(row=r_idx, column=c_idx, value=val)

        if cell_colors and sheet_name in cell_colors:
            for coord, hex_rgb in cell_colors[sheet_name].items():
                clean_rgb = hex_rgb.lstrip("#").upper()
                ws[coord].fill = openpyxl.styles.PatternFill(
                    start_color=clean_rgb, end_color=clean_rgb, fill_type="solid"
                )

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# 1. Pydantic Model Validation Tests
# ---------------------------------------------------------------------------

class TestMappingModels:

    def test_valid_template_creation(self):
        t = MappingTemplate(
            id="template-001",
            name="Q1 Production Template",
            description="Manufacturing line 1",
        )
        assert t.id == "template-001"
        assert t.version == 1
        assert t.blocks == []
        assert t.created_at is not None

    def test_invalid_template_id_rejected(self):
        with pytest.raises(ValueError, match="Invalid template ID"):
            MappingTemplate(id="../../etc/passwd", name="Bad ID")

    def test_column_role_validation(self):
        valid = ColumnRoleMapping(column_name="Defect", role="category")
        assert valid.role == "category"

        with pytest.raises(ValueError, match="Invalid column role"):
            ColumnRoleMapping(column_name="Defect", role="invalid_role")

    def test_ambiguous_row_decision_validation(self):
        d1 = AmbiguousRowDecision(source_row=5, sheet_row=6, decision="summary")
        assert d1.decision == "summary"

        with pytest.raises(ValueError, match="Invalid row decision"):
            AmbiguousRowDecision(source_row=5, sheet_row=6, decision="delete_now")

    def test_cell_colour_hex_validation(self):
        c1 = CellColourStatusMap(rgb="00b050", status_label="Completed")
        assert c1.rgb == "#00B050"

        c2 = CellColourStatusMap(rgb="#FF0000", status_label="Overdue")
        assert c2.rgb == "#FF0000"

        with pytest.raises(ValueError, match="Invalid RGB hex"):
            CellColourStatusMap(rgb="ZZ1234", status_label="Broken")

    def test_source_of_truth_confirmation(self):
        sot = SourceOfTruthConfirmation(
            confirmed=True,
            chosen_source="monthly_breakdown",
            confirmed_quantity=37111.0,
        )
        assert sot.confirmed is True
        assert sot.confirmed_quantity == 37111.0


# ---------------------------------------------------------------------------
# 2. Template Storage & Versioning Tests
# ---------------------------------------------------------------------------

class TestTemplateStorage:

    @pytest.fixture
    def temp_storage(self):
        temp_dir = Path(tempfile.mkdtemp())
        storage = TemplateStorage(storage_dir=temp_dir)
        yield storage
        shutil.rmtree(temp_dir, ignore_errors=True)

    def test_save_and_retrieve_template(self, temp_storage: TemplateStorage):
        t = MappingTemplate(id="t-001", name="Test Template")
        saved = temp_storage.save_template(t)
        assert saved.version == 1

        retrieved = temp_storage.get_template("t-001")
        assert retrieved is not None
        assert retrieved.name == "Test Template"
        assert retrieved.version == 1

    def test_version_increments_on_update(self, temp_storage: TemplateStorage):
        t = MappingTemplate(id="t-002", name="Version Test")
        temp_storage.save_template(t)
        assert t.version == 1

        # Save again with same ID
        t2 = MappingTemplate(id="t-002", name="Version Test Updated")
        saved2 = temp_storage.save_template(t2)
        assert saved2.version == 2
        assert saved2.name == "Version Test Updated"

        saved3 = temp_storage.save_template(saved2)
        assert saved3.version == 3

    def test_list_templates(self, temp_storage: TemplateStorage):
        t1 = MappingTemplate(id="t-1", name="Alpha")
        t2 = MappingTemplate(id="t-2", name="Beta")
        temp_storage.save_template(t1)
        temp_storage.save_template(t2)

        summaries = temp_storage.list_templates()
        assert len(summaries) == 2
        ids = [s.id for s in summaries]
        assert "t-1" in ids
        assert "t-2" in ids

    def test_delete_template(self, temp_storage: TemplateStorage):
        t = MappingTemplate(id="t-del", name="To Delete")
        temp_storage.save_template(t)
        assert temp_storage.get_template("t-del") is not None

        deleted = temp_storage.delete_template("t-del")
        assert deleted is True
        assert temp_storage.get_template("t-del") is None

        # Deleting non-existent returns False
        assert temp_storage.delete_template("t-del") is False

    def test_path_traversal_safe(self, temp_storage: TemplateStorage):
        # Path traversal characters should return None / False without raising or accessing parent dirs
        assert temp_storage.get_template("../../shadow") is None
        assert temp_storage.delete_template("../../shadow") is False


# ---------------------------------------------------------------------------
# 3. Normalization Engine Tests
# ---------------------------------------------------------------------------

class TestNormalizeWorkbook:

    def test_exclude_block_mapping(self):
        rows1 = [["Defect", "Count"], ["Scratch", 10], ["Dent", 5]]
        rows2 = [["Noise", "Ignored"], [1, 2], [3, 4]]
        raw = _make_xlsx({"Data": rows1, "Notes": rows2})

        template = MappingTemplate(
            id="t-ex",
            name="Exclusion Test",
            blocks=[
                BlockMapping(block_index=0, sheet="Data", range="A1:B3", included=True),
                BlockMapping(block_index=1, sheet="Notes", range="A1:B3", included=False),
            ],
        )

        result = normalize_workbook(raw, "test.xlsx", template)
        assert len(result.tables) == 1
        assert result.tables[0].sheet == "Data"

    def test_column_role_override(self):
        rows = [["Code", "Count"], ["1001", 10], ["1002", 20]]
        raw = _make_xlsx({"Data": rows})

        template = MappingTemplate(
            id="t-override",
            blocks=[
                BlockMapping(
                    block_index=0,
                    sheet="Data",
                    range="A1:B3",
                    column_roles={"Code": "category", "Count": "measure"},
                )
            ],
        )

        result = normalize_workbook(raw, "test.xlsx", template)
        assert len(result.tables) == 1
        table = result.tables[0]
        assert table.column_roles["Code"] == "category"
        assert table.column_roles["Count"] == "measure"

    def test_unpivot_wide_monthly_columns(self):
        rows = [
            ["Model", "Jan-26", "Feb-26"],
            ["Sedan", 100, 150],
            ["SUV",   200, 250],
        ]
        raw = _make_xlsx({"Production 2026": rows})

        template = MappingTemplate(
            id="t-unpivot",
            blocks=[
                BlockMapping(
                    block_index=0,
                    sheet="Production 2026",
                    range="A1:C3",
                    unpivot=UnpivotConfig(
                        enabled=True,
                        id_columns=["Model"],
                        value_columns=["Jan-26", "Feb-26"],
                        variable_name="Month",
                        value_name="Units",
                        measure_name="Production",
                    ),
                )
            ],
        )

        result = normalize_workbook(raw, "test.xlsx", template)
        assert len(result.tables) == 1
        table = result.tables[0]
        col_names = [c["name"] for c in table.columns]
        assert "Model" in col_names
        assert "Month" in col_names
        assert "Units" in col_names
        assert "Measure" in col_names
        # Long form rows: 2 models x 2 months = 4 rows
        assert table.row_count == 4

    def test_ambiguous_row_triage(self):
        rows = [
            ["Item", "Value"],
            ["Record 1", 10],
            ["Record 2", 20],
            ["Subtotal A", 30],  # row index 2 in df
            ["Record 3", 40],
        ]
        raw = _make_xlsx({"Data": rows})

        template = MappingTemplate(
            id="t-ambiguous",
            blocks=[
                BlockMapping(
                    block_index=0,
                    sheet="Data",
                    range="A1:B5",
                    ambiguous_row_decisions=[
                        AmbiguousRowDecision(source_row=2, sheet_row=4, decision="summary"),
                    ],
                )
            ],
        )

        result = normalize_workbook(raw, "test.xlsx", template)
        table = result.tables[0]
        # Clean records should not have Subtotal A
        record_items = [r[0] for r in table.rows]
        assert "Subtotal A" not in record_items
        assert len(table.summary_rows) == 1
        assert table.summary_rows[0][0] == "Subtotal A"

    def test_value_merge_replacement(self):
        rows = [
            ["Defect", "Count"],
            ["Scratch", 10],
            ["Scratchh", 5],  # Typo
            ["Dent", 8],
        ]
        raw = _make_xlsx({"Data": rows})

        template = MappingTemplate(
            id="t-merge",
            blocks=[
                BlockMapping(
                    block_index=0,
                    sheet="Data",
                    range="A1:B4",
                    value_merges=[
                        ValueMergeDecision(original_value="Scratchh", target_value="Scratch", accepted=True),
                        ValueMergeDecision(original_value="Dent", target_value="Dented", accepted=False),
                    ],
                )
            ],
        )

        result = normalize_workbook(raw, "test.xlsx", template)
        table = result.tables[0]
        defects = [r[0] for r in table.rows]
        assert "Scratchh" not in defects
        assert defects.count("Scratch") == 2
        # Rejected merge leaves Dent as Dent
        assert "Dent" in defects

    def test_cell_colour_status_mapping(self):
        rows = [
            ["Action", "Owner"],
            ["Replace drill bit", "John"],
            ["Calibrate sensor", "Alice"],
        ]
        # Color row 1 green (#00B050), row 2 red (#FF0000)
        colors = {
            "Actions": {
                "A2": "00B050",
                "A3": "FF0000",
            }
        }
        raw = _make_xlsx({"Actions": rows}, cell_colors=colors)

        template = MappingTemplate(
            id="t-colour",
            blocks=[
                BlockMapping(
                    block_index=0,
                    sheet="Actions",
                    range="A1:B3",
                    colour_mappings=[
                        CellColourStatusMap(rgb="#00B050", status_label="Completed"),
                        CellColourStatusMap(rgb="#FF0000", status_label="Overdue"),
                    ],
                )
            ],
        )

        result = normalize_workbook(raw, "test.xlsx", template)
        table = result.tables[0]
        col_names = [c["name"] for c in table.columns]
        assert "Action_Status" in col_names
        statuses = [r[-1] for r in table.rows]
        assert "Completed" in statuses
        assert "Overdue" in statuses

    def test_source_of_truth_confirmation_status(self):
        rows = [["Defect", "Count"], ["Scratch", 10]]
        raw = _make_xlsx({"Data": rows})

        # When no conflict is detected, ppm_ready is True
        t_ok = MappingTemplate(id="t-sot-ok")
        res_ok = normalize_workbook(raw, "test.xlsx", t_ok)
        assert res_ok.ppm_ready is True

        # When confirmed explicitly
        t_confirmed = MappingTemplate(
            id="t-sot-conf",
            source_of_truth=SourceOfTruthConfirmation(
                confirmed=True,
                chosen_source="monthly_breakdown",
                confirmed_quantity=37111.0,
            ),
        )
        res_conf = normalize_workbook(raw, "test.xlsx", t_confirmed)
        assert res_conf.ppm_ready is True


# ---------------------------------------------------------------------------
# 4. API Endpoint Integration Tests
# ---------------------------------------------------------------------------

class TestMappingAPIEndpoints:

    def test_save_and_get_mapping_template(self):
        import uuid
        t_id = f"api-test-{uuid.uuid4().hex[:8]}"
        payload = {
            "id": t_id,
            "name": "Integration Template",
            "description": "Created via API test",
            "blocks": [],
        }
        resp = client.post("/api/workbooks/mapping", json=payload)
        assert resp.status_code == 200
        data = resp.json()
        assert data["id"] == t_id
        assert data["version"] == 1

        # Get by ID
        get_resp = client.get(f"/api/workbooks/mapping/{t_id}")
        assert get_resp.status_code == 200
        assert get_resp.json()["name"] == "Integration Template"

        # Cleanup
        client.delete(f"/api/workbooks/mapping/{t_id}")

    def test_update_template_increments_version(self):
        import uuid
        t_id = f"api-test-v-{uuid.uuid4().hex[:8]}"
        payload1 = {"id": t_id, "name": "V1"}
        client.post("/api/workbooks/mapping", json=payload1)

        payload2 = {"id": t_id, "name": "V2"}
        resp2 = client.post("/api/workbooks/mapping", json=payload2)
        assert resp2.status_code == 200
        assert resp2.json()["version"] == 2

        # Cleanup
        client.delete(f"/api/workbooks/mapping/{t_id}")

    def test_list_templates_endpoint(self):
        resp = client.get("/api/workbooks/mappings")
        assert resp.status_code == 200
        assert isinstance(resp.json(), list)

    def test_get_nonexistent_template_returns_404(self):
        resp = client.get("/api/workbooks/mapping/nonexistent-xyz")
        assert resp.status_code == 404

    def test_delete_template_endpoint(self):
        import uuid
        t_id = f"api-del-{uuid.uuid4().hex[:8]}"
        client.post("/api/workbooks/mapping", json={"id": t_id, "name": "To Delete"})
        del_resp = client.delete(f"/api/workbooks/mapping/{t_id}")
        assert del_resp.status_code == 200
        assert del_resp.json()["status"] == "deleted"

        # Subsequent get returns 404
        assert client.get(f"/api/workbooks/mapping/{t_id}").status_code == 404

    def test_normalize_endpoint_with_inline_template(self):
        rows = [["Defect", "Count"], ["Scratch", 10], ["Dent", 5]]
        raw = _make_xlsx({"Sheet1": rows})

        template_dict = {
            "id": "inline-norm-01",
            "name": "Inline Norm",
            "blocks": [
                {
                    "block_index": 0,
                    "sheet": "Sheet1",
                    "range": "A1:B3",
                    "included": True,
                    "column_roles": {"Defect": "category", "Count": "measure"},
                }
            ],
        }

        files = {"file": ("test.xlsx", io.BytesIO(raw), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        data = {"template_json": json.dumps(template_dict)}

        resp = client.post("/api/workbooks/normalize", files=files, data=data)
        assert resp.status_code == 200
        body = resp.json()
        assert "tables" in body
        assert len(body["tables"]) >= 1
        assert body["tables"][0]["column_roles"]["Defect"] == "category"

    def test_normalize_endpoint_with_template_id(self):
        # Save template first
        t_id = "saved-norm-01"
        client.post(
            "/api/workbooks/mapping",
            json={
                "id": t_id,
                "name": "Saved Norm Template",
                "blocks": [],
            },
        )

        rows = [["A", "B"], [1, 2]]
        raw = _make_xlsx({"S": rows})
        files = {"file": ("s.xlsx", io.BytesIO(raw), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        data = {"template_id": t_id}

        resp = client.post("/api/workbooks/normalize", files=files, data=data)
        assert resp.status_code == 200
        assert resp.json()["template_id"] == t_id

    def test_normalize_endpoint_unauthenticated_returns_401(self):
        from app.auth import get_current_user
        app.dependency_overrides.pop(get_current_user, None)
        try:
            resp = client.post(
                "/api/workbooks/normalize",
                files={"file": ("test.csv", b"a,b\n1,2\n", "text/csv")},
            )
            assert resp.status_code == 401
        finally:
            from app.auth import AuthenticatedUser
            app.dependency_overrides[get_current_user] = lambda: AuthenticatedUser(
                id="00000000-0000-0000-0000-000000000001",
                email="testuser@example.com",
                role="authenticated",
            )
