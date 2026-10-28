"""
Milestone 6 Automated Test Suite — Report History, Export Engine & Security Hardening

Tests:
1. Dashboard saving, listing, reopening and recalculating from DuckDB Parquet data.
2. Dashboard renaming, duplicating, and deleting without deleting shared datasets.
3. Multi-tenant workspace isolation & RBAC permissions (viewer vs editor vs admin).
4. PDF export generation (ReportLab, %PDF header, layout, SOT warnings).
5. PNG export snapshot generation (Pillow, \x89PNG header, dimensions).
6. Excel workbook export (.xlsx with XlsxWriter, multiple sheets: Summary, Specs, Normalized Data).
7. Formula-injection prevention:
   - Untrusted strings starting with '=', '+', '-', '@' neutralized with leading quote.
   - Legitimate typed numbers (int, float, -150) remain numbers and are NOT escaped.
8. CSV export with UTF-8 BOM (utf-8-sig) and formula protection.
9. Source-of-truth reconciliation warning preservation across all exports.
10. Export lifecycle: Queueing via validated IDs, status polling, and authenticated download.
11. Cross-workspace export download rejection (HTTP 403).
12. Deletion of export records and storage cleanup.
"""

from __future__ import annotations

import io
import uuid
import openpyxl
import pandas as pd
import pytest
from fastapi.testclient import TestClient
from PIL import Image

from app.auth import AuthenticatedUser, get_current_user
from app.db.models import Dashboard, Dataset, Export, Upload, User, Workspace, WorkspaceMember
from app.db.repositories import DashboardRepository, DatasetRepository, ExportRepository, UploadRepository
from app.db.session import get_db
from app.exports.generator import (
    generate_csv_export,
    generate_excel_export,
    generate_export,
    generate_pdf_export,
    generate_png_export,
    prepare_dashboard_export_data,
    sanitize_for_spreadsheet,
)
from app.main import app
from app.storage import get_storage_backend, save_dataframe_as_parquet

client = TestClient(app)


# ---------------------------------------------------------------------------
# Test Fixtures & Synthetic Manufacturing Data
# ---------------------------------------------------------------------------

@pytest.fixture
def manufacturing_fixture():
    """
    Creates a verified synthetic manufacturing test dataset:
    - Monthly rejections & production counts
    - Defect Pareto breakdown
    - Action tracker countermeasure board
    - Persisted as Parquet in local storage and registered in DB
    """
    storage = get_storage_backend()
    db_gen = app.dependency_overrides[get_db]()
    db = next(db_gen)

    try:
        user_uuid = uuid.UUID("00000000-0000-0000-0000-000000000001")
        ws_uuid = uuid.UUID("00000000-0000-0000-0000-000000000001")

        # 1. Create Upload
        upload = Upload(
            id=uuid.uuid4(),
            user_id=user_uuid,
            workspace_id=ws_uuid,
            original_filename="Synthetic_Manufacturing_2026.xlsx",
            sanitized_filename="synthetic_manufacturing_2026.xlsx",
            storage_key="test_uploads/synthetic_mfg.xlsx",
            file_size_bytes=24500,
            file_type="xlsx",
            sha256_hash="0123456789abcdef0123456789abcdef0123456789abcdef0123456789abcdef",
            status="ready",
        )
        db.add(upload)
        db.commit()
        db.refresh(upload)

        # 2. Synthetic manufacturing records
        raw_df = pd.DataFrame([
            {"Month": "Jan-26", "Model": "157 TY", "Defect": "=cmd|' /C calc'!A0", "Count": 120, "Rejections": 120, "Action_Status": "open", "Champion": "@hacker", "Due_Date": "2026-02-15"},
            {"Month": "Jan-26", "Model": "367 TY", "Defect": "Flash / Burr", "Count": 85, "Rejections": 85, "Action_Status": "closed", "Champion": "Quality Lead", "Due_Date": "2026-01-20"},
            {"Month": "Feb-26", "Model": "157 TY", "Defect": "-danger_string", "Count": 65, "Rejections": 65, "Action_Status": "overdue", "Champion": "Prod Team", "Due_Date": "2026-02-01"},
            {"Month": "Feb-26", "Model": "457 TY", "Defect": "+formula_trick", "Count": 90, "Rejections": 90, "Action_Status": "open", "Champion": "+lead_dev", "Due_Date": "2026-03-01"},
            {"Month": "Mar-26", "Model": "Milwaukee", "Defect": "Dimensional Variation", "Count": 45, "Rejections": 45, "Action_Status": "closed", "Champion": "Tooling", "Due_Date": "2026-03-10"},
        ])

        parquet_key = f"datasets/{upload.id}/test_mfg.parquet"
        save_dataframe_as_parquet(raw_df, storage, parquet_key)

        # 3. Create Dataset record
        dataset = Dataset(
            id=uuid.uuid4(),
            upload_id=upload.id,
            sheet_name="Rejection_Summary",
            cell_range="A1:H6",
            kind="tidy",
            row_count=len(raw_df),
            column_count=len(raw_df.columns),
            parquet_storage_key=parquet_key,
            column_metadata=[{"name": col, "type": "string"} for col in raw_df.columns],
        )
        db.add(dataset)
        db.commit()
        db.refresh(dataset)

        # 4. Create Dashboard Specification with widgets
        spec = {
            "template_type": "manufacturing_quality",
            "version": 1,
            "title": "Manufacturing Executive Quality Review",
            "upload_id": str(upload.id),
            "source_of_truth_confirmed": False, # Conflict active
            "reconciliation": {
                "source_of_truth_confirmed": False,
                "has_sot_conflict": True,
                "conflict_details": "January production count conflict (37,111 vs 46,816)",
            },
            "available_filters": {
                "periods": ["Jan-26", "Feb-26", "Mar-26"],
                "models": ["157 TY", "367 TY", "457 TY", "Milwaukee"],
                "defects": ["Flash / Burr", "Dimensional Variation"],
            },
            "global_warnings": [
                "January 2026 production denominator conflict is unconfirmed."
            ],
            "widgets": [
                {
                    "widget_id": "widget_ppm_trend",
                    "widget_type": "ppm_trend_target",
                    "title": "Monthly Rejection PPM versus Target",
                    "dataset_id": str(dataset.id),
                    "dimensions": ["Month"],
                    "measures": ["Count"],
                    "aggregations": {"Count": "sum"},
                    "parameters": {"period_column": "Month", "measure_column": "Count", "target_ppm": 2000.0},
                    "provenance": {
                        "upload_id": str(upload.id),
                        "sheet_name": "Rejection_Summary",
                        "cell_range": "A1:H6",
                        "source_column": "Count",
                        "aggregation": "sum",
                        "confirmation_status": "unconfirmed_conflict",
                    },
                    "warnings": [],
                },
                {
                    "widget_id": "widget_pareto",
                    "widget_type": "defect_pareto",
                    "title": "Defect Pareto Breakdown",
                    "dataset_id": str(dataset.id),
                    "dimensions": ["Defect"],
                    "measures": ["Rejections"],
                    "aggregations": {"Rejections": "sum"},
                    "parameters": {"defect_column": "Defect", "measure_column": "Rejections", "pareto_threshold": 80.0},
                    "provenance": {
                        "upload_id": str(upload.id),
                        "sheet_name": "Rejection_Summary",
                        "cell_range": "A1:H6",
                        "source_column": "Rejections",
                        "aggregation": "sum",
                        "confirmation_status": "confirmed",
                    },
                    "warnings": [],
                },
                {
                    "widget_id": "widget_actions",
                    "widget_type": "action_tracker",
                    "title": "Corrective Action Tracking",
                    "dataset_id": str(dataset.id),
                    "dimensions": ["Champion"],
                    "measures": ["Count"],
                    "aggregations": {"Count": "count"},
                    "parameters": {
                        "status_column": "Action_Status",
                        "owner_column": "Champion",
                        "due_date_column": "Due_Date",
                    },
                    "provenance": {
                        "upload_id": str(upload.id),
                        "sheet_name": "Rejection_Summary",
                        "cell_range": "A1:H6",
                        "source_column": "Champion",
                        "aggregation": "count",
                        "confirmation_status": "confirmed",
                    },
                    "warnings": [],
                },
            ],
        }

        dash = Dashboard(
            id=uuid.uuid4(),
            upload_id=upload.id,
            workspace_id=ws_uuid,
            user_id=user_uuid,
            title="Manufacturing Executive Quality Review",
            spec=spec,
        )
        db.add(dash)
        db.commit()
        db.refresh(dash)

        return {
            "upload": upload,
            "dataset": dataset,
            "dashboard": dash,
            "raw_df": raw_df,
        }

    finally:
        db.close()


# ---------------------------------------------------------------------------
# Section 1: Report History & Dashboard CRUD Tests
# ---------------------------------------------------------------------------

def test_save_and_reopen_dashboard(manufacturing_fixture):
    """Test saving a dashboard spec and reopening it with DuckDB recalculation."""
    dash = manufacturing_fixture["dashboard"]

    # Reopen dashboard
    res = client.get(f"/api/dashboards/{dash.id}")
    assert res.status_code == 200
    data = res.json()
    assert data["id"] == str(dash.id)
    assert data["title"] == dash.title
    assert "rendered_widgets" in data
    assert len(data["rendered_widgets"]) >= 2

    # Check that widgets recalculated from Parquet
    pareto_w = next(w for w in data["rendered_widgets"] if w["widget_type"] == "defect_pareto")
    assert pareto_w["total_categories"] > 0
    assert len(pareto_w["items"]) > 0


def test_list_dashboards_with_source_info(manufacturing_fixture):
    """Test listing saved dashboards in workspace includes source filename."""
    res = client.get("/api/dashboards")
    assert res.status_code == 200
    dashboards = res.json()
    assert isinstance(dashboards, list)
    assert len(dashboards) >= 1
    d = dashboards[0]
    assert "title" in d
    assert "source_filename" in d
    assert d["source_filename"] == "Synthetic_Manufacturing_2026.xlsx"


def test_rename_dashboard(manufacturing_fixture):
    """Test renaming a dashboard via PATCH /api/dashboards/{id}."""
    dash = manufacturing_fixture["dashboard"]
    new_title = "Plant 1 - Q1 Quality Review Updated"

    res = client.patch(f"/api/dashboards/{dash.id}", json={"title": new_title})
    assert res.status_code == 200
    data = res.json()
    assert data["title"] == new_title

    # Verify persisted in GET
    get_res = client.get(f"/api/dashboards/{dash.id}")
    assert get_res.json()["title"] == new_title


def test_duplicate_dashboard(manufacturing_fixture):
    """Test duplicating dashboard creates an independent copy with new ID."""
    dash = manufacturing_fixture["dashboard"]

    res = client.post(f"/api/dashboards/{dash.id}/duplicate", json={"title": "Copy for Audit"})
    assert res.status_code == 201
    data = res.json()
    assert data["id"] != str(dash.id)
    assert data["title"] == "Copy for Audit"
    assert data["upload_id"] == str(dash.upload_id)


def test_delete_dashboard_preserves_shared_dataset(manufacturing_fixture):
    """
    CRITICAL: Deleting a dashboard must NEVER delete the shared source dataset
    or Parquet file referenced by other reports.
    """
    storage = get_storage_backend()
    dash = manufacturing_fixture["dashboard"]
    dataset = manufacturing_fixture["dataset"]

    # Verify Parquet file exists before delete
    assert storage.object_exists(dataset.parquet_storage_key)

    # Delete dashboard
    res = client.delete(f"/api/dashboards/{dash.id}")
    assert res.status_code == 200

    # Verify dashboard deleted
    get_res = client.get(f"/api/dashboards/{dash.id}")
    assert get_res.status_code == 404

    # VERIFY DATASET AND PARQUET REMAIN INTACT
    db_gen = app.dependency_overrides[get_db]()
    db = next(db_gen)
    try:
        ds_repo = DatasetRepository(db)
        retained_ds = ds_repo.get(dataset.id)
        assert retained_ds is not None, "Dataset was erroneously deleted when dashboard was deleted!"
        assert storage.object_exists(retained_ds.parquet_storage_key), "Parquet file was erroneously purged!"
    finally:
        db.close()


# ---------------------------------------------------------------------------
# Section 2: Security & Formula-Injection Mitigation Tests
# ---------------------------------------------------------------------------

def test_formula_injection_sanitization():
    """
    Verifies spreadsheet formula-injection mitigation:
    - Formula-triggering strings starting with '=', '+', '-', '@' prepended with single quote.
    - Pure numeric types (int, float) and valid numeric strings are NOT escaped.
    """
    # Malicious string inputs
    assert sanitize_for_spreadsheet("=SUM(A1:A10)") == "'=SUM(A1:A10)"
    assert sanitize_for_spreadsheet("=cmd|' /C calc'!A0") == "'=cmd|' /C calc'!A0"
    assert sanitize_for_spreadsheet("+cmd|' /C calc'!A0") == "'+cmd|' /C calc'!A0"
    assert sanitize_for_spreadsheet("-test_formula") == "'-test_formula"
    assert sanitize_for_spreadsheet("@external_link") == "'@external_link"

    # Legitimate numbers MUST NOT be converted or escaped
    assert sanitize_for_spreadsheet(-150) == -150
    assert sanitize_for_spreadsheet(42.5) == 42.5
    assert sanitize_for_spreadsheet(0) == 0
    assert sanitize_for_spreadsheet("-150.5") == "-150.5" # pure number string preserved
    assert sanitize_for_spreadsheet(None) is None
    assert sanitize_for_spreadsheet("Normal Text") == "Normal Text"


# ---------------------------------------------------------------------------
# Section 3: Export Formats (PDF, PNG, Excel, CSV) Tests
# ---------------------------------------------------------------------------

def test_pdf_export_generation(manufacturing_fixture):
    """Test generating a printable PDF report with ReportLab."""
    db_gen = app.dependency_overrides[get_db]()
    db = next(db_gen)
    dash = manufacturing_fixture["dashboard"]

    try:
        data = prepare_dashboard_export_data(db, str(dash.id))
        output_buffer = io.BytesIO()

        # Temporary file path for pdf
        import tempfile, os
        with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as f:
            tmp_path = f.name

        try:
            size = generate_pdf_export(data, tmp_path, {"orientation": "landscape"})
            assert size > 0
            with open(tmp_path, "rb") as f:
                pdf_bytes = f.read()

            # Verify PDF magic header
            assert pdf_bytes.startswith(b"%PDF-"), "Generated file is not a valid PDF!"
            assert len(pdf_bytes) > 1000
        finally:
            if os.path.exists(tmp_path):
                os.remove(tmp_path)
    finally:
        db.close()


def test_png_export_generation(manufacturing_fixture):
    """Test generating a high-DPI dashboard card snapshot with Pillow."""
    db_gen = app.dependency_overrides[get_db]()
    db = next(db_gen)
    dash = manufacturing_fixture["dashboard"]

    try:
        data = prepare_dashboard_export_data(db, str(dash.id))
        import tempfile, os
        with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as f:
            tmp_path = f.name

        try:
            size = generate_png_export(data, tmp_path)
            assert size > 0
            with open(tmp_path, "rb") as f:
                png_bytes = f.read()

            # Verify PNG magic header
            assert png_bytes.startswith(b"\x89PNG\r\n\x1a\n"), "Generated file is not a valid PNG!"

            # Open with Pillow and check dimensions
            img = Image.open(io.BytesIO(png_bytes))
            assert img.width == 1600
            assert img.height == 1100
        finally:
            if os.path.exists(tmp_path):
                os.remove(tmp_path)
    finally:
        db.close()


def test_excel_export_multi_sheet(manufacturing_fixture):
    """Test generating multi-sheet Excel workbook with XlsxWriter."""
    db_gen = app.dependency_overrides[get_db]()
    db = next(db_gen)
    dash = manufacturing_fixture["dashboard"]

    try:
        data = prepare_dashboard_export_data(db, str(dash.id))
        import tempfile, os
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            tmp_path = f.name

        try:
            size = generate_excel_export(data, tmp_path)
            assert size > 0

            # Verify with openpyxl
            wb = openpyxl.load_workbook(tmp_path, data_only=False)
            sheet_names = wb.sheetnames
            assert "Summary" in sheet_names
            assert "Dashboard Spec" in sheet_names
            assert any(s.startswith("Data_") for s in sheet_names)

            # Check formula injection sanitized in Data sheet
            data_sheet = next(wb[s] for s in sheet_names if s.startswith("Data_"))
            col_names = [data_sheet.cell(1, c).value for c in range(1, data_sheet.max_column + 1)]
            defect_col_idx = col_names.index("Defect") + 1

            # The malicious value '=cmd|' /C calc'!A0' must have been sanitized with single quote
            sanitized_cell = data_sheet.cell(2, defect_col_idx).value
            assert str(sanitized_cell).startswith("'="), f"Formula injection was not neutralized: {sanitized_cell}"

        finally:
            if os.path.exists(tmp_path):
                os.remove(tmp_path)
    finally:
        db.close()


def test_csv_export_utf8_bom(manufacturing_fixture):
    """Test generating CSV with UTF-8 BOM and formula protection."""
    db_gen = app.dependency_overrides[get_db]()
    db = next(db_gen)
    dash = manufacturing_fixture["dashboard"]

    try:
        data = prepare_dashboard_export_data(db, str(dash.id))
        import tempfile, os
        with tempfile.NamedTemporaryFile(suffix=".csv", delete=False) as f:
            tmp_path = f.name

        try:
            size = generate_csv_export(data, tmp_path, {"export_mode": "normalized"})
            assert size > 0
            with open(tmp_path, "rb") as f:
                raw_bytes = f.read()

            # Verify UTF-8 BOM
            assert raw_bytes.startswith(b"\xef\xbb\xbf"), "CSV missing UTF-8 BOM!"

            text_content = raw_bytes.decode("utf-8-sig")
            assert "Defect" in text_content
            # Verify formula escaping
            assert "'=cmd|" in text_content
        finally:
            if os.path.exists(tmp_path):
                os.remove(tmp_path)
    finally:
        db.close()


def test_sot_conflict_warning_preserved_in_exports(manufacturing_fixture):
    """Test that unresolved January production denominator conflict surfaces warning in export data."""
    db_gen = app.dependency_overrides[get_db]()
    db = next(db_gen)
    dash = manufacturing_fixture["dashboard"]

    try:
        data = prepare_dashboard_export_data(db, str(dash.id))
        assert not data["sot_confirmed"]
        assert data["conflict_warning"] is not None
        assert "37,111" in data["conflict_warning"]
        assert "46,816" in data["conflict_warning"]
    finally:
        db.close()


# ---------------------------------------------------------------------------
# Section 4: Export API, RBAC & End-to-End Download Tests
# ---------------------------------------------------------------------------

def test_export_api_queue_and_download(manufacturing_fixture):
    """Test queueing an export via API, verifying status and downloading."""
    dash = manufacturing_fixture["dashboard"]

    # 1. Queue export
    res = client.post(
        f"/api/dashboards/{dash.id}/exports",
        json={"format": "pdf", "options": {"orientation": "landscape"}},
    )
    assert res.status_code == 202
    data = res.json()
    assert "id" in data
    export_id = data["id"]
    assert data["format"] == "pdf"

    # 2. Check status
    status_res = client.get(f"/api/exports/{export_id}")
    assert status_res.status_code == 200
    status_data = status_res.json()
    assert status_data["id"] == export_id
    assert status_data["status"] in ("ready", "processing", "queued")

    # 3. Download generated export
    dl_res = client.get(f"/api/exports/{export_id}/download")
    assert dl_res.status_code == 200
    assert dl_res.headers["content-type"] == "application/pdf"
    assert "attachment" in dl_res.headers["content-disposition"]
    assert dl_res.content.startswith(b"%PDF-")


def test_viewer_role_restricted_from_raw_exports(manufacturing_fixture):
    """
    RBAC check: Viewers are permitted to export visual formats (PDF, PNG),
    but are strictly blocked (HTTP 403) from exporting raw datasets (Excel, CSV).
    """
    dash = manufacturing_fixture["dashboard"]

    # Temporarily switch user to viewer role
    def override_viewer_user():
        return AuthenticatedUser(
            id="00000000-0000-0000-0000-000000000001",
            email="viewer@example.com",
            display_name="Test Viewer",
            role="viewer",
            workspace_id="00000000-0000-0000-0000-000000000001",
            workspace_role="viewer",
        )

    app.dependency_overrides[get_current_user] = override_viewer_user

    try:
        # Viewer attempting XLSX export -> HTTP 403 FORBIDDEN
        res_xlsx = client.post(f"/api/dashboards/{dash.id}/exports", json={"format": "xlsx"})
        assert res_xlsx.status_code == 403
        assert "restricted to visual exports" in res_xlsx.json()["detail"]

        # Viewer attempting CSV export -> HTTP 403 FORBIDDEN
        res_csv = client.post(f"/api/dashboards/{dash.id}/exports", json={"format": "csv"})
        assert res_csv.status_code == 403

        # Viewer attempting PDF export -> HTTP 202 ACCEPTED (Permitted visual export)
        res_pdf = client.post(f"/api/dashboards/{dash.id}/exports", json={"format": "pdf"})
        assert res_pdf.status_code == 202

    finally:
        # Restore default admin override
        def override_admin_user():
            return AuthenticatedUser(
                id="00000000-0000-0000-0000-000000000001",
                email="testuser@example.com",
                display_name="Test User",
                role="admin",
                workspace_id="00000000-0000-0000-0000-000000000001",
                workspace_role="admin",
            )
        app.dependency_overrides[get_current_user] = override_admin_user


def test_export_format_validation():
    """Unsupported export format rejected with HTTP 422."""
    res = client.post(
        "/api/dashboards/00000000-0000-0000-0000-000000000001/exports",
        json={"format": "exe"},
    )
    assert res.status_code == 422
    assert "Unsupported export format" in res.json()["detail"]


def test_cross_workspace_export_download_rejection(manufacturing_fixture):
    """Cross-workspace isolation: user in Workspace B cannot access or download exports from Workspace A."""
    dash = manufacturing_fixture["dashboard"]

    # Create export in Workspace A
    res = client.post(f"/api/dashboards/{dash.id}/exports", json={"format": "png"})
    assert res.status_code == 202
    export_id = res.json()["id"]

    # Switch session to Workspace B
    def override_ws_b_user():
        return AuthenticatedUser(
            id="00000000-0000-0000-0000-000000000002",
            email="intruder@example.com",
            display_name="Other Workspace User",
            role="admin",
            workspace_id="00000000-0000-0000-0000-000000000002",
            workspace_role="admin",
        )

    app.dependency_overrides[get_current_user] = override_ws_b_user

    try:
        # Attempt status check -> 403
        status_res = client.get(f"/api/exports/{export_id}")
        assert status_res.status_code == 403

        # Attempt download -> 403
        dl_res = client.get(f"/api/exports/{export_id}/download")
        assert dl_res.status_code == 403
    finally:
        def override_admin_user():
            return AuthenticatedUser(
                id="00000000-0000-0000-0000-000000000001",
                email="testuser@example.com",
                display_name="Test User",
                role="admin",
                workspace_id="00000000-0000-0000-0000-000000000001",
                workspace_role="admin",
            )
        app.dependency_overrides[get_current_user] = override_admin_user


def test_delete_export_record_and_file(manufacturing_fixture):
    """Test deleting an export record purges the private storage object."""
    storage = get_storage_backend()
    dash = manufacturing_fixture["dashboard"]

    res = client.post(f"/api/dashboards/{dash.id}/exports", json={"format": "csv"})
    assert res.status_code == 202
    export_id = res.json()["id"]
    storage_key = res.json()["storage_key"]

    # Verify storage object exists
    assert storage.object_exists(storage_key)

    # Delete export
    del_res = client.delete(f"/api/exports/{export_id}")
    assert del_res.status_code == 200

    # Verify export record deleted
    get_res = client.get(f"/api/exports/{export_id}")
    assert get_res.status_code == 404

    # Verify storage object purged
    assert not storage.object_exists(storage_key)
